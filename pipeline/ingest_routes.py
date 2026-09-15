"""Ingest in-season routes run from a hand-supplied weekly export.

nflverse's only route source is the participation file, and FTN delivers that a season
at a time *after* the post-season (``seasons.PARTICIPATION``). So during the season the
route counts arrive as a spreadsheet supplied by hand, one per week, and this script
attaches them to the stat lines ``ingest_stats.py`` has already created.

**Charted routes, not pass-play participation.** 2016–2025 ``routes_run`` counts every
pass play a player was on the field for (``ingest_usage.py``), so a tight end kept in to
block still "ran a route". A charting export counts routes actually run. Within a season
the basis is uniform, and every percentile here is per season, so a board ranks fairly;
comparing a blocking tight end's route participation across seasons is not like for like.

**The files are never committed.** The export is licensed charting data and the repo is
public: ``data/routes/`` is gitignored apart from its README, and production is loaded
by running this script against it (see that README).

What it writes, per matched stat line:

  routes_run          = the export's count
  route_participation = routes / the team's **dropbacks** in that game, scrambles
                        included (play-by-play ``qb_dropback``)
  targets_per_route_run, yards_per_route_run
                      = re-derived from the stat line (``derive_route_rates``)

**Dropbacks, not ``ingest_usage.py``'s pass plays, and it was measured.** A charting
export counts routes on scrambles too — the receivers ran them; the quarterback just did
not throw. Against pass plays (which exclude scrambles), 10 of the 282 players in the
2026 Week 1 export came out above 100%, DeVonta Smith at 118%; against dropbacks, none
did, and no team's busiest receiver ever exceeded its dropbacks. Each basis keeps the
denominator its numerator was counted over.

Routes per game and fantasy points per route run are query-time registry metrics
(``routes_run_per_game``, ``fantasy_points_per_route_run``) and need nothing here.

The export's snaps, targets and carries are **cross-checked, never written**. The stored
values come from nflverse and PFR; a disagreement is logged for a human to look at.

Like every enrichment pass it only updates stat lines that exist, and it is idempotent.
A player who touched nothing has no nflverse stat line, so his routes have nowhere to go —
those are counted in the log rather than creating half-empty rows.
"""

import argparse
import csv
import logging
import re
from dataclasses import dataclass
from pathlib import Path

import nflreadpy as nfl
import polars as pl
from openpyxl import load_workbook
from sqlalchemy import text

from db import get_engine, load_team_id_map, upsert
from ingest_expert_boards import normalize_name
from ingest_usage import derive_route_rates

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
logger = logging.getLogger("pipeline.routes")

ROUTES_DIR = Path(__file__).parent / "data" / "routes"

# `<season>-w<week>[-label].xlsx` — e.g. 2026-w01.xlsx, 2026-w01-pre-mnf.xlsx. Files for
# the same week are applied in name order, so a later file wins where they overlap.
FILENAME_PATTERN = re.compile(
    r"^(?P<season>\d{4})-w(?P<week>\d{1,2})(?:-[a-z0-9-]+)?\.(?:xlsx|csv)$"
)

REQUIRED_COLUMNS = ("name", "position", "routes run", "team")
CROSS_CHECKS = {"snaps": "snap_count", "targets": "targets", "carries": "carries"}

# The export uses PFF's position labels. A fullback is a back for matching purposes.
POSITION_GROUPS = {"HB": "RB", "FB": "RB", "RB": "RB", "WR": "WR", "TE": "TE", "QB": "QB"}


@dataclass
class ExportRow:
    """One player-week from an export, with the file it came from."""

    name: str
    position: str
    team: str
    routes: int
    checks: dict[str, int | None]
    source: str


def _as_int(value) -> int | None:
    """An export cell as an integer, or None when it is blank."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    return int(float(value))


def read_export(path: Path) -> list[ExportRow]:
    """Read one export. Columns are found by header name, so their order does not matter."""
    if path.suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            table = [list(row) for row in csv.reader(handle)]
    else:
        workbook = load_workbook(path, read_only=True, data_only=True)
        table = [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]

    header = [str(cell or "").strip().lower() for cell in table[0]]
    missing = [column for column in REQUIRED_COLUMNS if column not in header]
    if missing:
        raise ValueError(f"{path.name}: missing column(s) {missing}; found {header}")
    index = {column: header.index(column) for column in header if column}

    rows = []
    for cells in table[1:]:
        if not any(cells):
            continue
        name = str(cells[index["name"]] or "").strip()
        routes = _as_int(cells[index["routes run"]])
        if not name or routes is None:
            logger.warning("%s: skipping a row with no name or no routes: %s", path.name, cells)
            continue
        rows.append(ExportRow(
            name=name,
            position=str(cells[index["position"]] or "").strip().upper(),
            team=str(cells[index["team"]] or "").strip(),
            routes=routes,
            checks={
                column: _as_int(cells[index[column]]) if column in index else None
                for column in CROSS_CHECKS
            },
            source=path.name,
        ))
    return rows


def _stat_lines(season: int, week: int) -> list[dict]:
    """Every stored stat line for one regular-season week, with the player's identity."""
    query = text(
        """
        SELECT s.player_id, s.game_id, s.team_id, p.name, p.position,
               s.snap_count, s.targets, s.carries
          FROM player_stats s
          JOIN players p USING (player_id)
         WHERE s.season = :season AND s.week = :week AND s.season_type = 'REG'
        """
    )
    with get_engine().connect() as connection:
        return [dict(row) for row in connection.execute(query, {"season": season, "week": week}).mappings()]


def _team_ids_by_name(season: int) -> dict[str, int]:
    """Full team name -> team_id, for the teams on that season's schedule.

    Restricted to the season because ``teams`` also holds historical franchises, and
    the export names a team the way it is called today.
    """
    query = text(
        """
        SELECT DISTINCT t.name, t.team_id
          FROM teams t
          JOIN games g ON t.team_id IN (g.home_team_id, g.away_team_id)
         WHERE g.season = :season
        """
    )
    with get_engine().connect() as connection:
        return {name.lower(): team_id for name, team_id in connection.execute(query, {"season": season})}


def _team_dropbacks(season: int, week: int) -> dict[tuple[str, int], int]:
    """``(game_id, team_id) -> dropbacks``, scrambles included — see the module docstring."""
    team_map = load_team_id_map()
    counts = (
        nfl.load_pbp([season])
        .filter(
            (pl.col("week") == week)
            & (pl.col("season_type") == "REG")
            & (pl.col("qb_dropback") == 1)
        )
        .group_by(["game_id", "posteam"])
        .agg(pl.len().alias("dropbacks"))
    )
    return {
        (row["game_id"], team_map[row["posteam"]]): row["dropbacks"]
        for row in counts.iter_rows(named=True)
        if row["posteam"] in team_map
    }


def _players_by_team() -> set[tuple[int, str]]:
    """``(team_id, normalised name)`` for every player we hold, to explain a miss."""
    with get_engine().connect() as connection:
        rows = connection.execute(text("SELECT team_id, name FROM players WHERE team_id IS NOT NULL"))
        return {(team_id, normalize_name(name)) for team_id, name in rows}


def _surname(folded: str) -> str:
    """The last token of a normalised name ("" for an empty one)."""
    return (folded.split() or [""])[-1]


def _match(rows: list[ExportRow], lines: list[dict], teams: dict[str, int]) -> tuple[dict, list, list]:
    """Pair export rows with stat lines on team plus normalised name.

    Falls back to a *unique* surname within the team and position group, which covers
    a nickname ("Bill" vs "Jacory"); every fallback is returned so it gets logged. Never
    guesses between two candidates.
    """
    exact: dict[tuple[int, str], list[dict]] = {}
    by_surname: dict[tuple[int, str, str], list[dict]] = {}
    for line in lines:
        folded = normalize_name(line["name"])
        exact.setdefault((line["team_id"], folded), []).append(line)
        group = POSITION_GROUPS.get(line["position"], line["position"])
        by_surname.setdefault((line["team_id"], _surname(folded), group), []).append(line)

    matched: dict[tuple[str, str], tuple[ExportRow, dict]] = {}
    fallbacks, unmatched = [], []
    for row in rows:
        team_id = teams.get(row.team.lower())
        if team_id is None:
            raise ValueError(f"{row.source}: unknown team {row.team!r} for {row.name}")
        folded = normalize_name(row.name)
        candidates = exact.get((team_id, folded), [])
        if len(candidates) != 1:
            group = POSITION_GROUPS.get(row.position, row.position)
            surname = by_surname.get((team_id, _surname(folded), group), [])
            if not candidates and len(surname) == 1:
                candidates = surname
                fallbacks.append((row, surname[0]))
        if len(candidates) != 1:
            unmatched.append(row)
            continue
        line = candidates[0]
        key = (line["player_id"], line["game_id"])
        if key in matched and matched[key][0].routes != row.routes:
            logger.warning(
                "%s appears twice for the same game (%s: %d routes, %s: %d) — the later file wins",
                line["name"], matched[key][0].source, matched[key][0].routes, row.source, row.routes,
            )
        matched[key] = (row, line)
    return matched, fallbacks, unmatched


def _report(
    matched: dict, fallbacks: list, unmatched: list, lines: list[dict],
    teams: dict[str, int], teams_in_file: set[int],
) -> None:
    """Log what matched, what did not, and where the export disagrees with stored stats."""
    logger.info("matched %d export rows to stat lines", len(matched))
    for row, line in fallbacks:
        logger.info("  matched on surname: %s (%s) -> %s", row.name, row.team, line["name"])

    no_line = [row for row in unmatched if not (row.checks["targets"] or row.checks["carries"])]
    if no_line:
        logger.info(
            "%d players touched nothing, so nflverse has no stat line to attach routes to: %s",
            len(no_line), ", ".join(f"{row.name} ({row.routes})" for row in no_line),
        )
    # A player the export credits with a target or carry but who matched no stat line is
    # one of two things, and they need different fixes: someone we hold but nflverse gave
    # nothing to this week, or someone we do not hold at all (a position outside scope,
    # or a name spelled differently).
    held = _players_by_team()
    for row in unmatched:
        if row in no_line:
            continue
        known = (teams[row.team.lower()], normalize_name(row.name)) in held
        logger.warning(
            "%s: %s, %s %s — %d routes, %s targets, %s carries",
            "NO STAT LINE (nflverse credits nothing this week)" if known
            else "UNMATCHED (not a player we hold — position outside scope, or a different spelling)",
            row.name, row.position, row.team, row.routes, row.checks["targets"], row.checks["carries"],
        )

    # The other direction: a stat line the export should have covered but did not.
    covered = set(matched)
    missed = [
        line for line in lines
        if line["team_id"] in teams_in_file
        and line["position"] in ("RB", "WR", "TE")
        and (line["targets"] or 0) > 0
        and (line["player_id"], line["game_id"]) not in covered
    ]
    for line in missed:
        logger.warning("NO ROUTES: %s has %d targets but no row in the export", line["name"], line["targets"])

    for column, stored in CROSS_CHECKS.items():
        compared = [(row, line) for row, line in matched.values() if row.checks[column] is not None]
        differ = [(row, line) for row, line in compared if row.checks[column] != (line[stored] or 0)]
        logger.info("cross-check %s: %d of %d agree", column, len(compared) - len(differ), len(compared))
        for row, line in differ:
            logger.warning(
                "  %s %s: export %s, stored %s", column, line["name"], row.checks[column], line[stored],
            )


def ingest_week(season: int, week: int, paths: list[Path], dry_run: bool = False) -> int:
    """Load every export for one week. Returns the stat lines written."""
    rows = [row for path in paths for row in read_export(path)]
    lines = _stat_lines(season, week)
    if not lines:
        logger.warning("%d week %d has no stat lines yet — run ingest_stats.py first", season, week)
        return 0

    teams = _team_ids_by_name(season)
    matched, fallbacks, unmatched = _match(rows, lines, teams)
    _report(matched, fallbacks, unmatched, lines, teams, {teams[row.team.lower()] for row in rows})

    dropbacks = _team_dropbacks(season, week)
    written_rows = []
    for (player_id, game_id), (row, line) in matched.items():
        denominator = dropbacks.get((game_id, line["team_id"]))
        participation = row.routes / denominator if denominator else None
        if participation is None:
            logger.warning("  %s: no dropback count for %s, so no route participation", line["name"], game_id)
        elif participation > 1:
            # Should not happen against dropbacks; if it does, the export and play-by-play
            # disagree about the game and a human should look before trusting either.
            logger.warning(
                "  %s: %d routes on %d team dropbacks (%.0f%%)",
                line["name"], row.routes, denominator, participation * 100,
            )
        written_rows.append({
            "player_id": player_id,
            "game_id": game_id,
            "routes_run": row.routes,
            "route_participation": participation,
        })

    if dry_run:
        logger.info("dry run: would write routes to %d stat lines", len(written_rows))
        return 0
    written = upsert("player_stats", written_rows, conflict_columns=["player_id", "game_id"])
    derive_route_rates([season])
    logger.info("routes %d week %d: wrote %d stat lines from %s", season, week, written, [p.name for p in paths])
    return written


def _discover() -> dict[tuple[int, int], list[Path]]:
    """Group the exports in data/routes by the week their filename names."""
    weeks: dict[tuple[int, int], list[Path]] = {}
    for path in sorted(ROUTES_DIR.glob("*")):
        if path.name == "README.md":
            continue
        match = FILENAME_PATTERN.match(path.name)
        if not match:
            logger.warning("skipping %s: expected a name like 2026-w01.xlsx", path.name)
            continue
        weeks.setdefault((int(match["season"]), int(match["week"])), []).append(path)
    return weeks


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Ingest hand-supplied weekly routes run.")
    parser.add_argument("--file", type=Path, help="One export, instead of everything in data/routes.")
    parser.add_argument("--season", type=int, help="Season for --file (else read from its name).")
    parser.add_argument("--week", type=int, help="Week for --file (else read from its name).")
    parser.add_argument("--dry-run", action="store_true", help="Match and cross-check without writing.")
    args = parser.parse_args()

    if args.file:
        named = FILENAME_PATTERN.match(args.file.name)
        season = args.season or (int(named["season"]) if named else None)
        week = args.week or (int(named["week"]) if named else None)
        if season is None or week is None:
            parser.error("--file needs --season and --week unless it is named like 2026-w01.xlsx")
        ingest_week(season, week, [args.file], dry_run=args.dry_run)
    else:
        discovered = _discover()
        if not discovered:
            logger.info("no route exports in %s", ROUTES_DIR)
        for (season, week), paths in sorted(discovered.items()):
            ingest_week(season, week, paths, dry_run=args.dry_run)

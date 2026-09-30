"""Ingest the season in progress's personnel groupings from a hand-supplied file.

nflverse's only personnel source is the participation file, which FTN delivers a season
at a time after the Super Bowl (``seasons.PARTICIPATION``), so ``ingest_team_stats.py``
fills ``team_personnel`` for 2016 to last season and nothing during one. For the season
in progress the numbers arrive by hand: one row per team, and for each grouping its
usage %, plays, EPA per play and success rate, season to date.

**Season totals, not games, so a table of their own.** The file says "through Week 3",
and a season total cannot be split back into games, so it is written to
``team_personnel_season`` rather than the per-game ``team_personnel``. The team pages
use it for any window that holds every game it counts (``app/team_stats.py``) and show
a dash for a narrower one rather than pretend.

**Stored as the file states it.** Usage share, EPA per play and success rate are kept
as given. The file's share is over every offensive play, including groupings it does
not list, so recomputing it from the listed plays would overstate every grouping for
a team that used an unlisted one (the Chargers' shares sum to 94%). The file withholds
EPA and success under 20 plays and writes "-"; those stay NULL.

**The files are never committed.** They are treated like the routes exports, since the
repo is public: ``data/personnel/`` is gitignored apart from its README.

Each file is a snapshot, so a newer one replaces the older team by team
(``replace_scoped``): a team missing from a newer file keeps its previous rows, with
their own ``through_week``, rather than losing them to a bad export.
"""

import argparse
import csv
import logging
import re
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

import franchises
from db import get_engine, load_team_id_map, replace_scoped

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
logger = logging.getLogger("pipeline.personnel")

PERSONNEL_DIR = Path(__file__).parent / "data" / "personnel"
# <season>-w<through week>[-label].xlsx, e.g. 2026-w03.xlsx
FILENAME_PATTERN = re.compile(r"^(?P<season>\d{4})-w(?P<week>\d{1,2})(?:-[\w-]+)?\.(?:xlsx|csv)$", re.IGNORECASE)
# "11 Personnel %", "11 Personnel Plays", "11 Personnel EPA", "11 Personnel Success Rate"
COLUMN_PATTERN = re.compile(r"^(?P<grouping>\d{2}) personnel (?P<field>%|plays|epa|success rate)$", re.IGNORECASE)
FIELDS = {"%": "share", "plays": "plays", "epa": "epa_per_play", "success rate": "success_rate"}


def _cell(value) -> float | None:
    """A number, or None for the file's blanks and "-"."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 6)  # Excel hands back 0.48100000000000004
    stripped = str(value).strip().rstrip("%")
    if stripped in ("", "-"):
        return None
    return round(float(stripped), 6)


def _read_rows(path: Path) -> list[dict]:
    """The file's rows as dicts keyed by header. ``.xlsx`` (first sheet) or ``.csv``."""
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            return list(csv.DictReader(handle))
    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    return [dict(zip(header, row)) for row in rows]


def read_export(path: Path, season: int, through_week: int, season_type: str = "REG") -> list[dict]:
    """Parse one file into ``team_personnel_season`` rows. Raises on anything ambiguous."""
    raw = _read_rows(path)
    if not raw:
        raise ValueError(f"{path.name}: empty file")
    columns = {}
    for header in raw[0]:
        match = COLUMN_PATTERN.match(header or "")
        if match:
            columns[header] = (match["grouping"], FIELDS[match["field"].lower()])
    if "Team" not in raw[0] or not columns:
        raise ValueError(f"{path.name}: expected a Team column and columns like '11 Personnel Plays'")

    team_ids = load_team_id_map()
    codes = franchises.contemporary_code_map([season])
    out: list[dict] = []
    seen: set[int] = set()
    ignored = 0
    for record in raw:
        team = record.get("Team")
        if not isinstance(team, str) or not team.strip():
            # stray cells below the table (the Week 3 file has a lone 100 there)
            ignored += any(v is not None and v != "" for v in record.values())
            continue
        code = franchises.resolve(codes, season, team.strip().upper())
        team_id = team_ids.get(code)
        if team_id is None:
            raise ValueError(f"{path.name}: unknown team {team!r}")
        if team_id in seen:
            raise ValueError(f"{path.name}: {team} appears twice")
        seen.add(team_id)

        groups: dict[str, dict] = {}
        for header, (grouping, field) in columns.items():
            groups.setdefault(grouping, {})[field] = _cell(record.get(header))
        for grouping, values in sorted(groups.items()):
            plays = values.get("plays") or 0
            if plays <= 0:
                continue
            for field, low, high in (("share", 0, 1), ("success_rate", 0, 1), ("epa_per_play", -3, 3)):
                value = values.get(field)
                if value is not None and not low <= value <= high:
                    raise ValueError(f"{path.name}: {team} {grouping} {field} = {value} is outside {low} to {high}")
            out.append({
                "team_id": team_id, "season": season, "season_type": season_type, "grouping": grouping,
                "through_week": through_week, "plays": plays, "share": values.get("share"),
                "epa_per_play": values.get("epa_per_play"), "success_rate": values.get("success_rate"),
            })
    if ignored:
        logger.info("%s: ignored %d row(s) with no team", path.name, ignored)
    _report(out, season, through_week, seen, team_ids)
    return out


def _report(rows: list[dict], season: int, through_week: int, seen: set[int], team_ids: dict[str, int]) -> None:
    """Log what a human should look at: missing teams, and plays against play-by-play."""
    with get_engine().connect() as connection:
        teams = {tid: abbr for abbr, tid in team_ids.items()}
        playing = {tid for (tid,) in connection.execute(text(
            "SELECT home_team_id FROM games WHERE season = :s AND week <= :w "
            "UNION SELECT away_team_id FROM games WHERE season = :s AND week <= :w"), {"s": season, "w": through_week})}
        pbp = dict(connection.execute(text(
            "SELECT team_id, SUM(plays) FROM team_game_stats WHERE season = :s AND week <= :w AND side = 'o' "
            "AND season_type = 'REG' GROUP BY team_id"), {"s": season, "w": through_week}).all())
    missing = sorted(teams[t] for t in playing - seen)
    if missing:
        logger.warning("teams not in the file (they keep any earlier rows): %s", ", ".join(missing))
    listed: dict[int, float] = {}
    for row in rows:
        listed[row["team_id"]] = listed.get(row["team_id"], 0) + row["plays"]
    logger.info("%d teams, %d grouping rows, through week %d", len(seen), len(rows), through_week)
    if not pbp:
        logger.warning("no play-by-play team stats through week %d yet, so plays are not cross-checked", through_week)
        return
    # The two never match exactly: the Week 3 file's listed groupings came to 86-98% of
    # play-by-play's offensive plays for every team (the low end is a team that used
    # groupings the file does not list). A file for the wrong week is off by a third.
    for team_id, plays in sorted(listed.items(), key=lambda kv: teams[kv[0]]):
        ours = pbp.get(team_id)
        if ours and abs(plays - ours) / ours > 0.2:
            logger.warning("  %s: the file lists %d plays in its groupings, play-by-play has %d", teams[team_id], plays, ours)


def ingest_file(path: Path, season: int, through_week: int, dry_run: bool = False) -> int:
    """Load one file. Returns rows written."""
    rows = read_export(path, season, through_week)
    if dry_run:
        logger.info("dry run: would write %d rows from %s", len(rows), path.name)
        return 0
    written = replace_scoped("team_personnel_season", rows, scope_columns=["team_id", "season", "season_type"])
    logger.info("personnel %d through week %d: wrote %d rows from %s", season, through_week, written, path.name)
    return written


def _newest_per_season() -> dict[int, tuple[int, Path]]:
    """The file with the latest through-week for each season in data/personnel."""
    newest: dict[int, tuple[int, Path]] = {}
    for path in sorted(PERSONNEL_DIR.glob("*")):
        if path.name == "README.md":
            continue
        match = FILENAME_PATTERN.match(path.name)
        if not match:
            logger.warning("skipping %s: expected a name like 2026-w03.xlsx", path.name)
            continue
        season, week = int(match["season"]), int(match["week"])
        if season not in newest or week >= newest[season][0]:
            newest[season] = (week, path)
    return newest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Ingest hand-supplied season-to-date personnel groupings.")
    parser.add_argument("--file", type=Path, help="One file, instead of the newest per season in data/personnel.")
    parser.add_argument("--season", type=int, help="Season for --file (else read from its name).")
    parser.add_argument("--week", type=int, help="The week --file runs through (else read from its name).")
    parser.add_argument("--dry-run", action="store_true", help="Parse and cross-check without writing.")
    args = parser.parse_args()

    if args.file:
        named = FILENAME_PATTERN.match(args.file.name)
        season = args.season or (int(named["season"]) if named else None)
        week = args.week or (int(named["week"]) if named else None)
        if season is None or week is None:
            parser.error("--file needs --season and --week unless it is named like 2026-w03.xlsx")
        ingest_file(args.file, season, week, dry_run=args.dry_run)
    else:
        found = _newest_per_season()
        if not found:
            logger.info("no personnel files in %s", PERSONNEL_DIR)
        for season, (week, path) in sorted(found.items()):
            ingest_file(path, season, week, dry_run=args.dry_run)
